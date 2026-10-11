"""Génération des livrables : rapport Word (note N2) et classeur Excel."""

import io
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from . import config as cfg
from .engine import AnalysisResult
from .narrative import build_findings, fmt_eur

NAVY = RGBColor(0x0D, 0x2A, 0x52)


def _add_table(doc, df: pd.DataFrame, font_size: int = 8):
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    for i, col in enumerate(df.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        for run in cell.paragraphs[0].runs:
            run.font.size, run.font.bold = Pt(font_size), True
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if isinstance(value, float):
                text = f"{value:,.2f}".replace(",", " ").replace(".", ",")
            elif isinstance(value, pd.Timestamp):
                text = value.strftime("%d/%m/%Y %H:%M")
            else:
                text = "" if pd.isna(value) else str(value)
            cells[i].text = text
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(font_size)
    return table


def _add_md(doc, text: str, style: str = None):
    """Paragraphe avec gras **...** converti en runs Word."""
    para = doc.add_paragraph(style=style)
    for i, chunk in enumerate(text.split("**")):
        run = para.add_run(chunk)
        run.bold = i % 2 == 1
    return para


def build_docx(res: AnalysisResult, source: str, notes: list = None) -> bytes:
    f = build_findings(res, source)
    df = res.df
    doc = Document()
    from docx.enum.section import WD_ORIENT
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = section.right_margin = Cm(1.5)

    title = doc.add_heading(f"{cfg.APP_NAME} — Note d'analyse transactionnelle", level=0)
    for run in title.runs:
        run.font.color.rgb = NAVY
    meta = doc.add_paragraph()
    meta.add_run(f"Source : {source}\n").bold = True
    meta.add_run(f"Date de l'analyse : {datetime.now():%d/%m/%Y %H:%M}\n"
                 f"Version de la plateforme : {cfg.VERSION}")

    doc.add_heading("1. Synthèse", level=1)
    doc.add_paragraph().add_run(f.headline).bold = True
    doc.add_paragraph(f.summary)
    kpis = pd.DataFrame([
        ("Opérations analysées", len(df)),
        ("Volume analysé", fmt_eur(df["MONTANT"].sum())),
        ("Opérations suspectes", len(f.table)),
        ("Alertes N1 (modéré)", int((df["NIVEAU"] == cfg.LEVEL_N1).sum())),
        ("Alertes N2 (élevé)", int((df["NIVEAU"] == cfg.LEVEL_N2).sum())),
    ], columns=["Indicateur", "Valeur"])
    _add_table(doc, kpis, 9)

    doc.add_heading("2. Ce que le moteur a relevé", level=1)
    if f.details:
        for text in f.details:
            _add_md(doc, text, style="List Bullet")
    else:
        doc.add_paragraph("Aucune typologie déclenchée.")
    doc.add_paragraph(f.clients_paragraph)

    doc.add_heading("3. Transactions suspectes", level=1)
    if len(f.table):
        _add_table(doc, f.table.head(60), 7)
        if len(f.table) > 60:
            doc.add_paragraph(f"Extrait des 60 opérations les plus élevées sur {len(f.table)} ; "
                              f"liste complète dans l'export Excel.")
    else:
        doc.add_paragraph("Aucune transaction suspecte.")

    doc.add_heading("4. Limites de l'analyse", level=1)
    doc.add_paragraph(f.limits)
    skipped = res.typology_status[res.typology_status["Statut"] != "Évaluée"]
    if len(skipped):
        _add_table(doc, skipped[["Code", "Typologie", "Colonnes manquantes"]])
    for note in notes or []:
        doc.add_paragraph(note, style="List Bullet")
    doc.add_paragraph(
        "Ce document est une aide à la décision générée automatiquement. Il ne constitue ni une "
        "qualification juridique ni une décision de déclaration : celle-ci relève de l'analyste "
        "N2 et du responsable LCB-FT / déclarant TRACFIN."
    ).runs[0].italic = True

    doc.add_heading("5. Conclusion de l'analyste N2", level=1)
    doc.add_paragraph("Décision : ☐ Classement sans suite   ☐ Surveillance renforcée   "
                      "☐ Demande d'informations complémentaires   ☐ Déclaration de soupçon")
    doc.add_paragraph("Commentaire : ______________________________________________")
    doc.add_paragraph("Analyste : ________________    Date : ____/____/______")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_excel_bytes(sheets: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
    return buf.getvalue()
